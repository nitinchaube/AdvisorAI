import os
import PyPDF2
import docx
import re
from typing import Dict, List, Any
from llm_service import LLMService
from dotenv import load_dotenv
import logging

load_dotenv()

# Set up logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class ResumeProcessor:
    def __init__(self):
        try:
            self.llm_service = LLMService()
            print(f"✅ LLM Service initialized with provider: {self.llm_service.provider}")
        except Exception as e:
            print(f"❌ LLM Service initialization failed: {e}")
            self.llm_service = None
    
    def extract_text_from_pdf(self, file_path: str) -> str:
        """Extract text from PDF file using multiple methods"""
        text = ""
        
        try:
            # Method 1: Try PyMuPDF (fitz) - better text extraction
            try:
                import fitz
                doc = fitz.open(file_path)
                for page_num in range(len(doc)):
                    page = doc.load_page(page_num)
                    page_text = page.get_text()
                    if page_text.strip():
                        text += page_text + "\n"
                doc.close()
                print(f"✅ PDF text extracted with PyMuPDF: {len(text)} characters")
                if len(text.strip()) > 50:
                    print(text)
                    return text
            except ImportError:
                print("⚠️ PyMuPDF not available, trying PyPDF2")
            except Exception as e:
                print(f"⚠️ PyMuPDF failed: {e}, trying PyPDF2")
            
            # Method 2: PyPDF2 as fallback
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    page_text = page.extract_text()
                    if page_text and page_text.strip():
                        text += page_text + "\n"
                
                print(f"✅ PDF text extracted with PyPDF2: {len(text)} characters")
                return text
                
        except Exception as e:
            raise Exception(f"Error extracting text from PDF: {str(e)}")
    
    def extract_text_from_docx(self, file_path: str) -> str:
        """Extract text from DOCX file"""
        try:
            doc = docx.Document(file_path)
            text = ""
            
            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text + "\n"
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text += cell.text + " "
                    text += "\n"
            
            print(f"✅ DOCX text extracted: {len(text)} characters")
            return text
            
        except Exception as e:
            raise Exception(f"Error extracting text from DOCX: {str(e)}")
    
    def extract_text_from_doc(self, file_path: str) -> str:
        """Extract text from DOC file"""
        try:
            # For DOC files, try to use antiword if available
            import subprocess
            try:
                result = subprocess.run(['antiword', file_path], capture_output=True, text=True)
                if result.returncode == 0:
                    text = result.stdout
                    print(f"✅ DOC text extracted with antiword: {len(text)} characters")
                    return text
            except (subprocess.SubprocessError, FileNotFoundError):
                print("⚠️ antiword not available, trying basic extraction")
            
            # Fallback: try to read as text
            try:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                    text = file.read()
                    print(f"✅ DOC text extracted with basic method: {len(text)} characters")
                    return text
            except UnicodeDecodeError:
                # Try different encodings
                for encoding in ['latin-1', 'cp1252', 'iso-8859-1']:
                    try:
                        with open(file_path, 'r', encoding=encoding, errors='ignore') as file:
                            text = file.read()
                            print(f"✅ DOC text extracted with {encoding}: {len(text)} characters")
                            return text
                    except UnicodeDecodeError:
                        continue
                
                raise Exception("Could not decode DOC file with any encoding")
                
        except Exception as e:
            raise Exception(f"Error extracting text from DOC: {str(e)}")
    
    def extract_text_from_file(self, file_path: str, file_type: str) -> str:
        """Extract text based on file type"""
        print(f"📄 Extracting text from {file_type} file: {file_path}")
        
        if file_type == 'application/pdf':
            return self.extract_text_from_pdf(file_path)
        elif file_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
            return self.extract_text_from_docx(file_path)
        elif file_type == 'application/msword':
            return self.extract_text_from_doc(file_path)
        else:
            raise Exception(f"Unsupported file type: {file_type}")
    
    def clean_text(self, text: str) -> str:
        """Clean and normalize extracted text"""
        if not text:
            return ""
        
        print(f"🧹 Cleaning text: {len(text)} characters")
        
        # Remove extra whitespace
        text = re.sub(r'\s+', ' ', text)
        
        # Remove special characters but keep basic punctuation
        text = re.sub(r'[^\w\s\.\,\;\:\!\?\-\(\)\@\#\$\&\+\=\[\]\{\}\|\:\"\'\\\/]', '', text)
        
        # Remove multiple periods, commas, etc.
        text = re.sub(r'\.{2,}', '.', text)
        text = re.sub(r'\,{2,}', ',', text)
        
        # Clean up common OCR artifacts
        text = re.sub(r'[|]{2,}', '|', text)
        text = re.sub(r'[-]{2,}', '-', text)
        
        cleaned_text = text.strip()
        print(f"✅ Text cleaned: {len(cleaned_text)} characters")
        
        return cleaned_text
    
    def parse_with_llm(self, text: str) -> Dict[str, Any]:
        """Parse resume text using configured LLM"""
        if not text or len(text.strip()) < 10:
            print("⚠️ Text too short for LLM parsing, using basic parsing")
            return self.basic_parse(text)
        
        if self.llm_service:
            return self.llm_service.parse_resume(text)
        else:
            # Fallback to basic parsing
            return self.basic_parse(text)
    
    def basic_parse(self, text: str) -> Dict[str, Any]:
        """Basic parsing without LLM"""
        print("🔍 Performing basic text parsing")
        
        # Extract email
        email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
        email_match = re.search(email_pattern, text)
        email = email_match.group() if email_match else ""
        
        # Extract phone
        phone_pattern = r'(\+\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}'
        phone_match = re.search(phone_pattern, text)
        phone = phone_match.group() if phone_match else ""
        
        # Extract name (basic pattern)
        name_pattern = r'^([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*)'
        name_match = re.search(name_pattern, text)
        full_name = name_match.group(1) if name_match else ""
        
        # Extract skills (basic keywords)
        skill_keywords = [
            'python', 'javascript', 'java', 'react', 'node.js', 'sql', 'mongodb',
            'aws', 'docker', 'kubernetes', 'git', 'html', 'css', 'typescript',
            'angular', 'vue', 'php', 'ruby', 'go', 'rust', 'c++', 'c#', '.net',
            'machine learning', 'ai', 'data science', 'devops', 'cloud computing',
            'html5', 'css3', 'jquery', 'bootstrap', 'express', 'django', 'flask',
            'postgresql', 'mysql', 'redis', 'elasticsearch', 'kafka', 'rabbitmq'
        ]
        skills = [skill for skill in skill_keywords if skill.lower() in text.lower()]
        
        print(f"✅ Basic parsing completed - Email: {email}, Phone: {phone}, Skills: {len(skills)}")
        
        return {
            "fullName": full_name,
            "email": email,
            "phone": phone,
            "location": "",
            "summary": "",
            "experience": [],
            "education": [],
            "skills": skills,
            "certifications": [],
            "projects": []
        }
    
    def process_resume(self, file_path: str, file_type: str) -> Dict[str, Any]:
        """Main method to process resume file"""
        try:
            print(f"🚀 Starting resume processing for: {file_path}")
            
            # Extract text
            text = self.extract_text_from_file(file_path, file_type)
            
            if not text or len(text.strip()) < 10:
                return {
                    "success": False,
                    "error": "No text could be extracted from the file",
                    "parsedData": {},
                    "originalText": ""
                }
            
            print(f"📝 Extracted text length: {len(text)} characters")
            print(f"📝 Text preview: {text[:200]}...")
            
            # Clean text
            cleaned_text = self.clean_text(text)
            
            if not cleaned_text or len(cleaned_text.strip()) < 10:
                return {
                    "success": False,
                    "error": "Text cleaning resulted in empty content",
                    "parsedData": {},
                    "originalText": text
                }
            
            # Parse with LLM
            parsed_data = self.parse_with_llm(cleaned_text)
            
            print(f"✅ Resume processing completed successfully")
            print(f"📊 Parsed data keys: {list(parsed_data.keys())}")
            
            return {
                "success": True,
                "parsedData": parsed_data,
                "originalText": cleaned_text,
                "llmProvider": self.llm_service.provider if self.llm_service else "none"
            }
            
        except Exception as e:
            print(f"❌ Resume processing error: {str(e)}")
            return {
                "success": False,
                "error": str(e),
                "parsedData": {},
                "originalText": ""
            } 